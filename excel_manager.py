import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, date, time
import json

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "construction_expenses.xlsx")
METADATA_FILE = os.path.join(BASE_DIR, "records_site_meta.json")
PROJECTS_FILE = os.path.join(BASE_DIR, "projects.json")

HEADERS = [
    "Expense ID",
    "Site ID",
    "Site Name",
    "Date",
    "Material Name",
    "Quantity",
    "Unit",
    "Material Cost (INR)",
    "Labour/Work Type",
    "Number of Workers",
    "Labour Cost (INR)",
    "Payment Mode",
    "Brief Description",
    "Total Expense (INR)",
    "Site Photo"
]

def safe_float(val):
    if val is None or val == "" or val == "-":
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    if isinstance(val, datetime):
        excel_start = datetime(1899, 12, 30)
        return float((val - excel_start).days)
    if isinstance(val, date):
        excel_start = date(1899, 12, 30)
        return float((val - excel_start).days)
    if isinstance(val, time):
        return 0.0
    try:
        cleaned = str(val).replace("Rs.", "").replace("₹", "").replace("INR", "").replace(",", "").replace("$", "").strip()
        if not cleaned or cleaned == "-":
            return 0.0
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0

def safe_int(val):
    if val is None or val == "" or val == "-":
        return 0
    if isinstance(val, int):
        return val
    if isinstance(val, float):
        return int(val)
    if isinstance(val, (datetime, date)):
        return int(safe_float(val))
    try:
        cleaned = str(val).replace(",", "").strip()
        if not cleaned or cleaned == "-":
            return 0
        return int(float(cleaned))
    except (ValueError, TypeError):
        return 0

def safe_str(val):
    if val is None:
        return ""
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    return str(val).strip()

def load_projects_data():
    if os.path.exists(PROJECTS_FILE):
        try:
            with open(PROJECTS_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def load_records_meta():
    if os.path.exists(METADATA_FILE):
        try:
            with open(METADATA_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}

def save_records_meta(meta):
    try:
        with open(METADATA_FILE, "w", encoding="utf-8") as f:
            json.dump(meta, f, indent=2)
    except Exception as e:
        print("Warning: Could not save records metadata:", e)

def get_header_styles():
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='medium', color='0F172A')
    )
    return header_fill, header_font, header_align, thin_border

def get_data_styles(is_even=False):
    fill_color = "F8FAFC" if is_even else "FFFFFF"
    data_fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    data_font = Font(name="Calibri", size=10)
    data_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    return data_fill, data_font, data_border

def setup_sheet_headers(ws):
    ws.views.sheetView[0].showGridLines = True
    header_fill, header_font, header_align, thin_border = get_header_styles()
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

def adjust_column_widths(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = safe_str(cell.value)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def init_excel():
    if not os.path.exists(EXCEL_FILE):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "All_Projects_Summary"
        setup_sheet_headers(ws)
        wb.save(EXCEL_FILE)

    sync_project_sheets()

def sync_project_sheets():
    """Ensure every site project has its own dedicated sheet tab in the Excel workbook."""
    if not os.path.exists(EXCEL_FILE):
        return

    wb = openpyxl.load_workbook(EXCEL_FILE)
    projects = load_projects_data()
    all_records = get_all_records_raw(wb)

    if "All_Projects_Summary" not in wb.sheetnames:
        ws_main = wb.create_sheet(title="All_Projects_Summary", index=0)
        setup_sheet_headers(ws_main)
    else:
        ws_main = wb["All_Projects_Summary"]
        setup_sheet_headers(ws_main)

    # Remove any sheets for projects that have been deleted
    valid_sheets = {"All_Projects_Summary"} | {safe_str(sid).upper() for sid in projects.keys()}
    for s_name in list(wb.sheetnames):
        if s_name not in valid_sheets:
            try:
                del wb[s_name]
            except Exception:
                pass

    for site_id, p_info in projects.items():
        sheet_title = safe_str(site_id).upper()
        if sheet_title not in wb.sheetnames:
            ws_site = wb.create_sheet(title=sheet_title)
            setup_sheet_headers(ws_site)
        else:
            ws_site = wb[sheet_title]
            setup_sheet_headers(ws_site)

        while ws_site.max_row > 1:
            ws_site.delete_rows(2)

        site_records = [r for r in all_records if safe_str(r.get("Site ID")).upper() == sheet_title]
        for row_idx, r in enumerate(site_records, start=2):
            is_even = (row_idx % 2 == 0)
            data_fill, data_font, data_border = get_data_styles(is_even)
            row_vals = [
                (r["Expense ID"], "center", "@"),
                (r["Site ID"], "center", "@"),
                (r["Site Name"], "left", None),
                (r["Date"], "center", "yyyy-mm-dd"),
                (r["Material Name"], "left", None),
                (r["Quantity"] if r["Quantity"] > 0 else "-", "right", "#,##0.00" if r["Quantity"] > 0 else None),
                (r["Unit"] if r["Unit"] else "-", "center", None),
                (r["Material Cost (INR)"], "right", "#,##0.00"),
                (r["Labour/Work Type"], "left", None),
                (r["Number of Workers"] if r["Number of Workers"] > 0 else "-", "right", "#,##0" if r["Number of Workers"] > 0 else None),
                (r["Labour Cost (INR)"], "right", "#,##0.00"),
                (r.get("Payment Mode", "Cash"), "center", None),
                (r["Brief Description"], "left", None),
                (r["Total Expense (INR)"], "right", "#,##0.00"),
                (r["Site Photo"] if r["Site Photo"] else "-", "center", None)
            ]
            for col_i, (val, align_h, num_fmt) in enumerate(row_vals, start=1):
                cell = ws_site.cell(row=row_idx, column=col_i, value=val)
                cell.fill = data_fill
                cell.font = data_font
                cell.border = data_border
                cell.alignment = Alignment(horizontal=align_h, vertical="center")
                if num_fmt:
                    cell.number_format = num_fmt
            ws_site.row_dimensions[row_idx].height = 24
        adjust_column_widths(ws_site)

    wb.save(EXCEL_FILE)

def get_next_id(site_id=None):
    rows = get_all_records(site_filter=site_id if site_id and site_id.upper() != "ALL" else None)
    if not rows:
        return "EXP-001"
    max_num = 0
    for r in rows:
        exp_id = safe_str(r.get("Expense ID", ""))
        clean_id = exp_id.upper()
        if "EXP-" in clean_id:
            try:
                part = clean_id.split("EXP-")[-1].strip()
                num = int(part)
                if num > max_num:
                    max_num = num
            except ValueError:
                pass
    return f"EXP-{max_num + 1:03d}"

def add_record(data):
    init_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb.active if "All_Projects_Summary" not in wb.sheetnames else wb["All_Projects_Summary"]

    next_row = ws.max_row + 1
    is_even = (next_row % 2 == 0)
    data_fill, data_font, data_border = get_data_styles(is_even)

    site_id = safe_str(data.get("Site ID") or data.get("site_id") or data.get("siteId") or "SITE-101")
    expense_id = safe_str(data.get("Expense ID") or data.get("expenseId")) or get_next_id(site_id)
    site_name = safe_str(data.get("Site Name") or data.get("site_name") or data.get("siteName") or "Villa Grand Residency")
    date_val = safe_str(data.get("Date") or data.get("entryDate")) or datetime.now().strftime("%Y-%m-%d")
    material_name = safe_str(data.get("Material Name") or data.get("materialName"))
    quantity = safe_float(data.get("Quantity") or data.get("quantity"))
    unit = safe_str(data.get("Unit") or data.get("unit"))
    
    mat_cost_val = (
        data.get("Material Cost (INR)")
        or data.get("Material Cost")
        or data.get("materialCost")
        or data.get("material_cost")
        or data.get("Material Cost (₹)")
    )
    material_cost = safe_float(mat_cost_val)
    
    work_type = safe_str(data.get("Labour/Work Type") or data.get("workType")) or "General Site Work"
    num_workers = safe_int(data.get("Number of Workers") or data.get("numWorkers"))
    
    lab_cost_val = (
        data.get("Labour Cost (INR)")
        or data.get("Labour Cost")
        or data.get("labourCost")
        or data.get("labour_cost")
        or data.get("Labour Cost (₹)")
    )
    labour_cost = safe_float(lab_cost_val)
    
    payment_mode = safe_str(data.get("Payment Mode") or data.get("payment_mode") or data.get("paymentMode") or "Cash")
    if payment_mode not in ["Cash", "Online Payment", "Cheque"]:
        payment_mode = "Cash"

    description = safe_str(data.get("Brief Description") or data.get("briefDesc"))
    
    total_val = data.get("Total Expense (INR)") or data.get("total_expense")
    total_expense = safe_float(total_val)
    if total_expense == 0.0 or total_expense is None:
        total_expense = material_cost + labour_cost

    photo_url = safe_str(data.get("Site Photo") or data.get("photo_url") or data.get("photo") or "")

    row_values = [
        (expense_id, "center", "@"),
        (site_id, "center", "@"),
        (site_name, "left", None),
        (date_val, "center", "yyyy-mm-dd"),
        (material_name, "left", None),
        (quantity if quantity > 0 else "-", "right", "#,##0.00" if quantity > 0 else None),
        (unit if unit else "-", "center", None),
        (material_cost, "right", "#,##0.00"),
        (work_type, "left", None),
        (num_workers if num_workers > 0 else "-", "right", "#,##0" if num_workers > 0 else None),
        (labour_cost, "right", "#,##0.00"),
        (payment_mode, "center", None),
        (description, "left", None),
        (total_expense, "right", "#,##0.00"),
        (photo_url if photo_url else "-", "center", None)
    ]

    for col_idx, (val, align_h, num_fmt) in enumerate(row_values, start=1):
        cell = ws.cell(row=next_row, column=col_idx, value=val)
        cell.fill = data_fill
        cell.font = data_font
        cell.border = data_border
        cell.alignment = Alignment(horizontal=align_h, vertical="center")
        if num_fmt:
            cell.number_format = num_fmt

    ws.row_dimensions[next_row].height = 24
    adjust_column_widths(ws)
    wb.save(EXCEL_FILE)

    meta = load_records_meta()
    meta[expense_id] = {
        "expense_id": expense_id,
        "site_id": site_id,
        "site_name": site_name,
        "photo_url": photo_url,
        "date": date_val,
        "payment_mode": payment_mode,
        "total_expense": total_expense
    }
    save_records_meta(meta)

    sync_project_sheets()

    return {
        "Expense ID": expense_id,
        "Site ID": site_id,
        "Site Name": site_name,
        "Date": date_val,
        "Material Name": material_name,
        "Quantity": quantity,
        "Unit": unit,
        "Material Cost (INR)": material_cost,
        "Labour/Work Type": work_type,
        "Number of Workers": num_workers,
        "Labour Cost (INR)": labour_cost,
        "Payment Mode": payment_mode,
        "Brief Description": description,
        "Total Expense (INR)": total_expense,
        "Site Photo": photo_url
    }

add_entry = add_record

def get_all_records_raw(wb):
    ws = wb.active if "All_Projects_Summary" not in wb.sheetnames else wb["All_Projects_Summary"]
    records_meta = load_records_meta()
    records = []
    
    header_map = {}
    for col_idx in range(1, ws.max_column + 1):
        h_val = safe_str(ws.cell(row=1, column=col_idx).value).lower().strip()
        if h_val:
            header_map[h_val] = col_idx

    for row_idx in range(2, ws.max_row + 1):
        c1 = ws.cell(row=row_idx, column=1).value
        if c1 is None or safe_str(c1) == "":
            continue

        exp_id = safe_str(c1)
        m_info = records_meta.get(exp_id, {})

        if header_map:
            site_id = safe_str(ws.cell(row=row_idx, column=header_map.get("site id", 2)).value) if "site id" in header_map else m_info.get("site_id", "SITE-101")
            site_name = safe_str(ws.cell(row=row_idx, column=header_map.get("site name", 3)).value) if "site name" in header_map else m_info.get("site_name", "Villa Grand Residency")
            date_str = safe_str(ws.cell(row=row_idx, column=header_map.get("date", 4)).value)
            mat_name = safe_str(ws.cell(row=row_idx, column=header_map.get("material name", 5)).value)
            qty = safe_float(ws.cell(row=row_idx, column=header_map.get("quantity", 6)).value)
            unit = safe_str(ws.cell(row=row_idx, column=header_map.get("unit", 7)).value)
            mat_cost = safe_float(ws.cell(row=row_idx, column=header_map.get("material cost (inr)", 8)).value)
            work_type = safe_str(ws.cell(row=row_idx, column=header_map.get("labour/work type", 9)).value)
            num_workers = safe_int(ws.cell(row=row_idx, column=header_map.get("number of workers", 10)).value)
            lab_cost = safe_float(ws.cell(row=row_idx, column=header_map.get("labour cost (inr)", 11)).value)
            
            if "payment mode" in header_map:
                pay_mode = safe_str(ws.cell(row=row_idx, column=header_map["payment mode"]).value) or m_info.get("payment_mode", "Cash")
            else:
                pay_mode = m_info.get("payment_mode", "Cash")

            desc = safe_str(ws.cell(row=row_idx, column=header_map.get("brief description", 13 if "payment mode" in header_map else 12)).value)
            tot_val = ws.cell(row=row_idx, column=header_map.get("total expense (inr)", 14 if "payment mode" in header_map else 13)).value
            tot_exp = safe_float(tot_val)
            photo_val = safe_str(ws.cell(row=row_idx, column=header_map.get("site photo", 15 if "payment mode" in header_map else 14)).value)
        else:
            site_id = m_info.get("site_id", "SITE-101")
            site_name = m_info.get("site_name", "Villa Grand Residency")
            date_str = safe_str(ws.cell(row=row_idx, column=2).value)
            mat_name = safe_str(ws.cell(row=row_idx, column=3).value)
            qty = safe_float(ws.cell(row=row_idx, column=4).value)
            unit = safe_str(ws.cell(row=row_idx, column=5).value)
            mat_cost = safe_float(ws.cell(row=row_idx, column=6).value)
            work_type = safe_str(ws.cell(row=row_idx, column=7).value)
            num_workers = safe_int(ws.cell(row=row_idx, column=8).value)
            lab_cost = safe_float(ws.cell(row=row_idx, column=9).value)
            pay_mode = m_info.get("payment_mode", "Cash")
            desc = safe_str(ws.cell(row=row_idx, column=10).value)
            tot_val = ws.cell(row=row_idx, column=11).value
            tot_exp = safe_float(tot_val)
            photo_val = safe_str(ws.cell(row=row_idx, column=12).value)

        if not pay_mode:
            pay_mode = "Cash"
        if tot_exp == 0.0:
            tot_exp = mat_cost + lab_cost
        if not photo_val or photo_val == "-":
            photo_val = m_info.get("photo_url", "")

        records.append({
            "Expense ID": exp_id,
            "Site ID": site_id,
            "Site Name": site_name,
            "Date": date_str,
            "Material Name": mat_name,
            "Quantity": qty,
            "Unit": unit,
            "Material Cost (INR)": mat_cost,
            "Labour/Work Type": work_type,
            "Number of Workers": num_workers,
            "Labour Cost (INR)": lab_cost,
            "Payment Mode": pay_mode,
            "Brief Description": desc,
            "Total Expense (INR)": tot_exp,
            "Site Photo": photo_val
        })
    return records

def get_all_records(site_filter=None):
    if not os.path.exists(EXCEL_FILE):
        return []
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True)
    records = get_all_records_raw(wb)

    if site_filter and site_filter.upper() != "ALL":
        records = [r for r in records if r["Site ID"].upper() == site_filter.upper() or r["Site Name"].upper() == site_filter.upper()]

    return records

def create_project_excel_file(site_id, project_info, records):
    """Creates a dedicated standalone Excel file for a single client project."""
    file_path = os.path.join(BASE_DIR, f"construction_expenses_{site_id}.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{site_id}_Project_Report"
    ws.views.sheetView[0].showGridLines = True

    ws.merge_cells("A1:O1")
    title_cell = ws["A1"]
    title_cell.value = f"BUILDTRACK PRO - CONSTRUCTION PROJECT REPORT ({site_id})"
    title_cell.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    header_fill, header_font, header_align, thin_border = get_header_styles()
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_align
        cell.border = thin_border
    ws.row_dimensions[2].height = 28

    for row_idx, r in enumerate(records, start=3):
        is_even = (row_idx % 2 == 0)
        data_fill, data_font, data_border = get_data_styles(is_even)
        row_vals = [
            (r["Expense ID"], "center", "@"),
            (r["Site ID"], "center", "@"),
            (r["Site Name"], "left", None),
            (r["Date"], "center", "yyyy-mm-dd"),
            (r["Material Name"], "left", None),
            (r["Quantity"] if r["Quantity"] > 0 else "-", "right", "#,##0.00" if r["Quantity"] > 0 else None),
            (r["Unit"] if r["Unit"] else "-", "center", None),
            (r["Material Cost (INR)"], "right", "#,##0.00"),
            (r["Labour/Work Type"], "left", None),
            (r["Number of Workers"] if r["Number of Workers"] > 0 else "-", "right", "#,##0" if r["Number of Workers"] > 0 else None),
            (r["Labour Cost (INR)"], "right", "#,##0.00"),
            (r.get("Payment Mode", "Cash"), "center", None),
            (r["Brief Description"], "left", None),
            (r["Total Expense (INR)"], "right", "#,##0.00"),
            (r["Site Photo"] if r["Site Photo"] else "-", "center", None)
        ]
        for col_i, (val, align_h, num_fmt) in enumerate(row_vals, start=1):
            cell = ws.cell(row=row_idx, column=col_i, value=val)
            cell.fill = data_fill
            cell.font = data_font
            cell.border = data_border
            cell.alignment = Alignment(horizontal=align_h, vertical="center")
            if num_fmt:
                cell.number_format = num_fmt
        ws.row_dimensions[row_idx].height = 24

    adjust_column_widths(ws)
    wb.save(file_path)
    return file_path

def delete_record(expense_id):
    if not os.path.exists(EXCEL_FILE):
        return False

    target_id_str = safe_str(expense_id).upper().strip()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    
    deleted_any = False
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for row_idx in range(ws.max_row, 1, -1):
            cell_val = safe_str(ws.cell(row=row_idx, column=1).value).upper().strip()
            if cell_val == target_id_str:
                ws.delete_rows(row_idx)
                deleted_any = True

    wb.save(EXCEL_FILE)

    meta = load_records_meta()
    keys_to_del = [k for k in list(meta.keys()) if k.upper().strip() == target_id_str]
    for k in keys_to_del:
        del meta[k]
        deleted_any = True
    save_records_meta(meta)

    sync_project_sheets()

    for fname in os.listdir(BASE_DIR):
        if fname.startswith("construction_expenses_") and fname.endswith(".xlsx") and fname != "construction_expenses.xlsx":
            try:
                os.remove(os.path.join(BASE_DIR, fname))
            except Exception:
                pass

    return deleted_any

def delete_site_data(site_id):
    if not os.path.exists(EXCEL_FILE):
        return False
    site_id_str = safe_str(site_id).upper().strip()
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)

        # 1. Remove project-specific sheet if present
        if site_id_str in wb.sheetnames:
            del wb[site_id_str]

        # 2. Delete rows belonging to this site from All_Projects_Summary
        if "All_Projects_Summary" in wb.sheetnames:
            ws_main = wb["All_Projects_Summary"]
            for row_idx in range(ws_main.max_row, 1, -1):
                cell_val = safe_str(ws_main.cell(row=row_idx, column=2).value).upper().strip()
                if cell_val == site_id_str:
                    ws_main.delete_rows(row_idx)

        wb.save(EXCEL_FILE)
    except Exception as e:
        print(f"Error removing site {site_id_str} from Excel:", e)

    # 3. Clean records metadata
    try:
        meta = load_records_meta()
        keys_to_del = [k for k, v in meta.items() if safe_str(v.get("site_id")).upper().strip() == site_id_str]
        for k in keys_to_del:
            del meta[k]
        save_records_meta(meta)
    except Exception:
        pass

    # 4. Remove generated project report file
    proj_report = os.path.join(BASE_DIR, f"construction_expenses_{site_id_str}.xlsx")
    if os.path.exists(proj_report):
        try:
            os.remove(proj_report)
        except Exception:
            pass

    sync_project_sheets()
    return True